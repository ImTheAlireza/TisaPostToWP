"""فاز ۷: مبدل فایل ردیابی ۲.۰ — سؤال به‌جای حدس، لینک به‌جای «ردیف ۴»، و فایلِ برگشتی.

Run with ``python -m pytest tests/test_tracking.py`` or ``python3 -m unittest discover -s tests``.

What is pinned here is the part a warehouse notices:

* :mod:`bot.services.processor` — a file whose headers are ambiguous is **asked about**,
  never guessed (a guessed column is how a wrong code reaches the tracking system), the
  row limits actually stop the run, and a barcode Excel destroyed is refused — including
  when it comes back as text after being pasted into a cell;
* ``needs-review.xlsx`` as a *form*, not a report: its barcode column is text and its header
  row is the shop's own wording, so the same file goes back into the bot and comes out as
  a correct ``tracking.csv``;
* :mod:`bot.services.tracking_ledger` — a repeat of the same export is recognised and
  re-run only on purpose;
* :mod:`bot.modules.tracking_converter` — the buttons that ask the questions are real
  ``ConversationHandler`` entry points, and the download they wait on is released.

Files here are built with the real writers (openpyxl / csv / pymupdf), because the bugs
this module had were exactly the ones a fake in-memory frame cannot show: a float cell,
a cp1256 byte, a span at the wrong x.
"""

from __future__ import annotations

import csv
import io
import os
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

os.environ.setdefault("BOT_TOKEN", "123456:TEST")
os.environ.setdefault("SUDO_IDS", "1234567")

import _flow_harness as h

try:
    from bot.services import barcodes, processor

    HAS_PROCESSOR = True
except Exception:  # pragma: no cover - pandas missing
    barcodes = processor = None  # type: ignore[assignment]
    HAS_PROCESSOR = False

try:
    import openpyxl  # noqa: F401  (the xlsx fixtures are written with it)

    HAS_OPENPYXL = True
except Exception:  # pragma: no cover
    HAS_OPENPYXL = False

try:
    import pymupdf

    HAS_PYMUPDF = True
except Exception:  # pragma: no cover
    pymupdf = None  # type: ignore[assignment]
    HAS_PYMUPDF = False

try:
    from bot.constants import CB
    from bot.modules import tracking_converter as TC
    from bot.services import tracking_ledger

    HAS_FLOW = True
except Exception:  # pragma: no cover - PTB missing
    CB = TC = tracking_ledger = None  # type: ignore[assignment]
    HAS_FLOW = False

needs_processor = unittest.skipUnless(HAS_PROCESSOR, "pandas is not installed")
needs_xlsx = unittest.skipUnless(HAS_PROCESSOR and HAS_OPENPYXL, "pandas/openpyxl are not installed")
needs_pdf = unittest.skipUnless(HAS_PROCESSOR and HAS_PYMUPDF, "pandas/pymupdf are not installed")
needs_flow = unittest.skipUnless(HAS_PROCESSOR and HAS_FLOW, "pandas or python-telegram-bot is not installed")

SUDO = 1234567
HEADER = ["ردیف", "بارکد", "تاریخ ثبت", "نام گیرنده", "کد سفارش", "مقصد"]
GOOD = "610001573845123456789012"  # 24 digits: the Tisa rule, no checksum
GOOD2 = "610001573845123456789013"
EAN_OK = "4006381333931"  # a published valid EAN-13
EAN_BAD = "4006381333932"  # the same, last digit wrong


# ---------------------------------------------------------------------------
# Real files, written the way the shop's own tools write them
# ---------------------------------------------------------------------------


class FileFixture(unittest.TestCase):
    """A temp directory plus the three writers, so a test says what a file holds."""

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.dir = Path(self._tmp.name)
        # A suite that wrote into the repo's real data/ could hand the next run a
        # "already processed" answer for a file it never saw.
        self._ledger = h.temp_ledger()
        self._ledger.__enter__()
        self.addCleanup(self._ledger.__exit__, None, None, None)

    def xlsx(self, rows: list[list[object]], name: str = "orders.xlsx") -> Path:
        from openpyxl import Workbook

        path = self.dir / name
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        for row in rows:
            ws.append(row)
        wb.save(path)
        return path

    def csv(self, rows: list[list[object]], name: str = "orders.csv", encoding: str = "cp1256") -> Path:
        """A CSV written the way a Persian Windows tool writes one.

        cp1256 has no U+06CC (Persian yeh), so the fixture spells its headers with the
        Arabic variants — the same bytes a real export carries, and the reason
        ``_norm`` maps ي→ی and ك→ک.
        """
        import csv as _csv

        path = self.dir / name
        with open(path, "w", encoding=encoding, newline="") as handle:
            writer = _csv.writer(handle, delimiter=",", lineterminator="\r\n")
            writer.writerows(rows)
        return path

    def pdf(self, rows: list[tuple[str, str, str]], name: str = "orders.pdf") -> Path:
        """The Tisa PDF layout: code at the right, barcode in the middle, text left.

        ``processor._read_pdf`` reads fixed x bands (that is how the printed form is
        laid out), so the fixture places text there on purpose — and the matching test
        checks that a *moved* layout is refused loudly, not answered with an empty CSV.
        """
        path = self.dir / name
        doc = pymupdf.open()
        page = doc.new_page(width=595, height=842)
        y = 90
        for number, barcode, middle in rows:
            page.insert_text((560, y), number, fontsize=9)
            page.insert_text((430, y), barcode, fontsize=9)
            page.insert_text((240, y), middle, fontsize=9)
            y += 16
        doc.save(path)
        doc.close()
        return path


def csv_rows(report: object) -> list[str]:
    return report.csv_text.strip().splitlines()  # type: ignore[attr-defined]


def problems(report: object) -> list[dict[str, str]]:
    """`problems.csv` read back — the file the warehouse actually opens, so the test
    asserts on that and not on an in-memory tuple that may never reach the disk."""
    text = report.problems_csv or ""  # type: ignore[attr-defined]
    return [dict(row) for row in csv.DictReader(io.StringIO(text))]


def answers(seen: list) -> list[str]:
    """The toast/alert texts a query was answered with, in order."""
    return [str(entry[1] or "") for entry in seen if entry[0] == "answer"]


# ---------------------------------------------------------------------------
# 7.1 the column map: ask, never guess
# ---------------------------------------------------------------------------


@needs_xlsx
class TestColumnMap(FileFixture):
    def test_one_clear_header_needs_no_question(self):
        path = self.xlsx([HEADER, [1, GOOD, "1403-01-01", "امیر ۱۲۳۴۵", "123456", "تهران"]])
        report = processor.process_file(path, path.name)
        self.assertFalse(report.needs_answer)
        self.assertEqual(csv_rows(report)[1:], ["123456," + GOOD])
        self.assertIn("بارکد ← B (هدر)", report.summary)
        self.assertIn("کد سفارش ← E (هدر)", report.summary)

    def test_two_barcode_headers_become_a_question_not_a_guess(self):
        path = self.xlsx(
            [
                ["ردیف", "بارکد", "کد رهگیری", "کد سفارش"],
                [1, GOOD, EAN_OK, "123456"],
                [2, GOOD2, EAN_OK, "654321"],
            ]
        )
        report = processor.process_file(path, path.name)
        self.assertTrue(report.needs_answer)
        self.assertEqual(report.csv_text, "")  # nothing was written yet
        self.assertEqual(len(report.questions), 1)
        question = report.questions[0]
        self.assertEqual(question.field, processor.FIELD_BARCODE)
        labels = [option.label() for option in question.options]
        self.assertEqual(len(labels), 2)
        # the button says which column it is, and shows a value from it
        self.assertTrue(any("بارکد" in label and "B" in label for label in labels), labels)
        self.assertTrue(any("کد رهگیری" in label and "C" in label for label in labels), labels)

    def test_answering_the_question_reads_the_chosen_column(self):
        path = self.xlsx(
            [
                ["ردیف", "بارکد", "کد رهگیری", "کد سفارش"],
                [1, GOOD, GOOD2, "123456"],
            ]
        )
        first = processor.process_file(path, path.name)
        chosen = next(o.index for o in first.questions[0].options if o.header == "کد رهگیری")
        layout = processor.Layout.from_dict({**first.layout.as_dict(), "barcode_col": chosen})
        report = processor.process_file(path, path.name, layout=layout)
        self.assertFalse(report.needs_answer)
        self.assertEqual(csv_rows(report)[1:], ["123456," + GOOD2])
        self.assertIn("بارکد ← C (انتخاب تو)", report.summary)

    def test_a_file_without_headers_is_asked_and_loses_no_row(self):
        # No header row at all: the old code either guessed or skipped the first data
        # row as a label. Asking is the honest answer, and after the answer row 1 is
        # still data.
        path = self.xlsx(
            [
                [1, GOOD, "1403-01-01", "امیر ۱۲۳۴۵", "123456", "تهران"],
                [2, GOOD2, "1403-01-02", "رضا ۶۵۴۳۲", "654321", "مشهد"],
            ]
        )
        report = processor.process_file(path, path.name)
        self.assertTrue(report.needs_answer)
        self.assertIn("بارکد", report.questions[0].prompt)
        options = report.questions[0].options
        self.assertEqual(len(options), 6)  # every column is on offer
        self.assertTrue(any("61000157" in option.sample for option in options), [o.sample for o in options])
        layout = processor.Layout.from_dict({"header_row": -1, "barcode_col": 1, "code_col": 4})
        answered = processor.process_file(path, path.name, layout=layout)
        self.assertEqual(answered.rows, 2)
        self.assertEqual(len(csv_rows(answered)) - 1, 2)

    def test_the_skip_answer_says_so_in_the_report(self):
        path = self.xlsx(
            [
                ["ردیف", "بارکد", "نام گيرنده", "کد سفارش", "کد سفارش گیرنده"],
                [1, GOOD, "امیر ۱۲۳۴۵", "123456", "654321"],
            ]
        )
        report = processor.process_file(path, path.name)
        question = next(q for q in report.questions if q.field == processor.FIELD_CODE)
        self.assertTrue(question.allow_skip)
        skipped = processor.process_file(
            path, path.name, layout=report.layout.with_choice(processor.FIELD_CODE, -1)
        )
        self.assertFalse(skipped.needs_answer)
        # an explicit «no order-code column» also closes the name-column fallback:
        # reading a code the user just refused would be a different answer than theirs
        self.assertEqual(skipped.csv_text.strip().splitlines()[1], "," + GOOD)
        # the summary says the column is off *and* that the user said so, not that it
        # was missing — the difference matters when the file is re-read next week
        self.assertIn("کد سفارش ← ندارد (انتخاب تو)", skipped.summary)
        self.assertIn("حتی از نام گیرنده", skipped.summary)
        self.assertIn("کد سفارش خالی", skipped.problems_csv or "")

        # a sheet that simply has no code column is another story: the code is lifted
        # out of the name, and the summary admits exactly that
        without = self.xlsx(
            [["ردیف", "بارکد", "نام گيرنده"], [1, GOOD, "امیر ۱۲۳۴۵"], [2, GOOD2, "رضا ۶۵۴۳۲"]]
        )
        lifted = processor.process_file(without, without.name)
        self.assertEqual(
            csv_rows(lifted)[1:], ["12345," + GOOD, "65432," + GOOD2]
        )
        self.assertIn("2 کد سفارش از ستونِ نام گیرنده برداشته شد", lifted.summary)

    def test_an_answered_but_useless_choice_says_so_instead_of_returning_nothing(self):
        # Nothing in this sheet looks like a barcode, so the bot asks; and once it has
        # been pointed at a column of words it must not answer with an empty CSV — it
        # says the file had no data rows.
        path = self.xlsx([["الف", "ب"], ["یک", "دو"], ["سه", "چهار"]])
        report = processor.process_file(path, path.name)
        self.assertTrue(report.needs_answer)
        layout = processor.Layout.from_dict({"header_row": -1, "barcode_col": 0, "code_col": 1})
        with self.assertRaises(ValueError) as caught:
            processor.process_file(path, path.name, layout=layout)
        self.assertIn("هیچ ردیف داده‌ای در فایل پیدا نشد", str(caught.exception))


# ---------------------------------------------------------------------------
# 7.2 barcodes: length rules from settings, checksums where they exist
# ---------------------------------------------------------------------------


@needs_processor
class TestBarcodes(unittest.TestCase):
    def test_check_digit_matches_the_published_examples(self):
        # Each of these is a real, published code; if the algorithm drifts, this fails.
        for code in (
            "4006381333931",
            "5901234123457",
            "036000291452",
            "042100005264",
            "96385074",
            "12345670",
        ):
            self.assertEqual(barcodes.gtin_check_digit(code[:-1]), int(code[-1]), code)

    def test_classify(self) -> None:
        self.assertEqual(barcodes.classify(GOOD), ("ok", ""))
        self.assertEqual(barcodes.classify(EAN_OK)[0], "warn")
        state, note = barcodes.classify(EAN_BAD)
        self.assertEqual(state, "error")
        self.assertIn("رقم کنترلی", note)
        self.assertIn("1", note)  # the digit it should have ended with
        self.assertEqual(barcodes.classify("1234567890123456789")[0], "error")
        self.assertEqual(barcodes.classify("")[0], "error")
        self.assertEqual(barcodes.classify("abc123")[0], "error")


@needs_xlsx
class TestBarcodePolicy(FileFixture):
    def test_a_valid_ean13_is_kept_and_named_in_the_report(self):
        path = self.xlsx([HEADER, [1, EAN_OK, "1403-01-01", "امیر ۱۲۳۴۵", "123456", "x"]])
        report = processor.process_file(path, path.name)
        self.assertEqual(csv_rows(report)[1:], ["123456," + EAN_OK])
        self.assertIn("بارکد EAN-13 است، نه کد رهگیری تیسا", report.problems_csv or "")
        self.assertEqual(report.usable, 1)

    def test_an_ean13_with_a_broken_check_digit_never_reaches_the_csv(self):
        path = self.xlsx([HEADER, [1, EAN_BAD, "1403-01-01", "امیر ۱۲۳۴۵", "123456", "x"]])
        report = processor.process_file(path, path.name)
        self.assertEqual(len(csv_rows(report)), 1)  # header only
        self.assertIn("رقم کنترلی", report.review_csv or "")
        self.assertEqual(report.usable, 0)


# ---------------------------------------------------------------------------
# 7.3 limits: a big file is refused, not half-processed
# ---------------------------------------------------------------------------


@needs_xlsx
class TestDuplicates(FileFixture):
    """A repeated barcode is a *warning*: one parcel with two orders is real, and
    deleting both rows is how a correct order disappears. But the report still has to
    say every place it appears, in a form a human can read."""

    def test_a_repeated_barcode_is_kept_and_located(self) -> None:
        rows = [[index, EAN_OK, "1403-01-01", "امیر ۱۲۳۴۵", f"12345{index % 6}", "x"] for index in range(1, 6)]
        path = self.xlsx([HEADER, *rows])
        report = processor.process_file(path, path.name)
        self.assertEqual(len(csv_rows(report)) - 1, 5)  # nothing was dropped
        line = next(p for p in problems(report) if "بارکد تکراری" in p["reason"])
        self.assertIn("5 سطر", line["reason"])
        self.assertEqual(line["severity"], "⚠️")
        self.assertEqual(line["source_cell"], "Sheet1!B2 · Sheet1!B3 · Sheet1!B4")
        self.assertIn("+2 سطر دیگر", line["source_row"])  # three names, then the count

    def test_a_repeated_order_code_is_reported_once_per_code(self) -> None:
        path = self.xlsx(
            [
                HEADER,
                [1, GOOD, "1403-01-01", "امیر ۱۲۳۴۵", "123456", "x"],
                [2, GOOD2, "1403-01-02", "رضا ۶۵۴۳۲", "123456", "y"],
            ]
        )
        report = processor.process_file(path, path.name)
        repeats = [p for p in problems(report) if "کد سفارش تکراری" in p["reason"]]
        self.assertEqual(len(repeats), 1, problems(report))
        self.assertEqual(repeats[0]["source_cell"], "Sheet1!E2 · Sheet1!E3")


@needs_xlsx
class TestLimits(FileFixture):
    def test_too_many_rows_is_refused_with_the_limit_named(self):
        rows = [HEADER] + [[i, GOOD, "1403-01-01", f"گیرنده {i}", "123456", "x"] for i in range(1, 6)]
        path = self.xlsx(rows)
        small = h.settings_with(max_rows=3)
        with h.patched_settings(small), self.assertRaises(processor.RowLimitError) as caught:
            processor.process_file(path, path.name)
        message = str(caught.exception)
        self.assertIn("MAX_ROWS", message)
        self.assertIn("6", message)  # what the file holds

    def test_the_same_file_is_fine_under_a_bigger_limit(self):
        rows = [HEADER] + [[i, GOOD, "1403-01-01", f"گیرنده {i}", "123456", "x"] for i in range(1, 6)]
        path = self.xlsx(rows)
        with h.patched_settings(h.settings_with(max_rows=100)):
            report = processor.process_file(path, path.name)
        self.assertEqual(report.rows, 5)


# ---------------------------------------------------------------------------
# 7.4 problems.csv with a row link, needs-fix as a form
# ---------------------------------------------------------------------------


@needs_xlsx
class TestFixFiles(FileFixture):
    def test_problems_csv_points_at_the_cell_that_is_wrong(self):
        path = self.xlsx(
            [
                HEADER,
                [1, GOOD, "1403-01-01", "امیر ۱۲۳۴۵", "123456", "x"],
                [2, 1.93e23, "1403-01-02", "رضا ۶۵۴۳۲", "654321", "y"],  # Excel destroyed it
                [3, "6100015738451234567", "1403-01-03", "هدیه", "77889", "z"],
                ["جمع کل", "", "", "", "", ""],
            ]
        )
        report = processor.process_file(path, path.name)
        problems = report.problems_csv or ""
        self.assertTrue(problems.startswith("severity,source_row,source_cell,reason,order_id,barcode"))
        # the destroyed barcode is in B3 and the 5-digit code in E4 — sheet rows, not
        # "the third data row", so «جمع کل» cannot shift the address
        self.assertIn("Sheet1!B3", problems)
        self.assertIn("Sheet1!E4", problems)
        self.assertNotIn("Sheet1!B5", problems)  # the total row is not data

    def test_a_plain_csv_source_gets_a_row_number_not_a_fake_cell(self):
        rows = [
            ["رديف", "باركد", "تاريخ ثبت", "نام گيرنده", "كد سفارش", "مقصد"],
            [1, "6100015738451234567", "1403-01-01", "هديه", "77889", "z"],
        ]
        path = self.csv(rows)
        report = processor.process_file(path, path.name)
        self.assertIn("سطر 2 فایل", report.problems_csv or "")
        self.assertNotIn("Sheet1!", report.problems_csv or "")

    def test_the_fix_workbook_is_a_form_the_bot_can_read_back(self):
        """The round trip: blank where nothing can be recovered, text everywhere else."""
        path = self.xlsx(
            [
                HEADER,
                [1, 1.93e23, "1403-01-01", "امیر ۱۲۳۴۵", "123456", "x"],
                [2, "6100015738451234567", "1403-01-02", "رضا ۶۵۴۳۲", "654321", "y"],
            ]
        )
        first = processor.process_file(path, path.name)
        self.assertIsNotNone(first.review_xlsx)
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(first.review_xlsx))  # type: ignore[arg-type]
        ws = wb["needs-review"]
        self.assertEqual([c.value for c in ws[1]], ["ردیف", "بارکد", "کد سفارش", "دلیل", "محل در فایل اصلی"])
        # a destroyed number is not offered back as if it were data
        self.assertIn(ws.cell(row=2, column=2).value, (None, ""))
        # the 19-digit one is, so it can be fixed in place — and it is TEXT
        self.assertEqual(ws.cell(row=3, column=2).value, "6100015738451234567")
        self.assertEqual(ws.cell(row=3, column=2).number_format, "@")
        self.assertIn("بارکد در اکسل خراب شده", str(ws.cell(row=2, column=4).value))

        ws.cell(row=2, column=2).value = GOOD  # what a human does in Excel
        ws.cell(row=3, column=2).value = GOOD2
        fixed = self.dir / "needs-review.xlsx"
        wb.save(fixed)

        again = processor.process_file(fixed, fixed.name)
        self.assertFalse(again.needs_answer)
        self.assertEqual(csv_rows(again)[1:], ["123456," + GOOD, "654321," + GOOD2])
        self.assertIsNone(again.review_csv)
        self.assertIn("🎉 هیچ مشکلی پیدا نشد!", again.summary)


# ---------------------------------------------------------------------------
# 7.5 the PDF path: the printed form, and a layout change that is admitted
# ---------------------------------------------------------------------------


@needs_pdf
class TestPdf(FileFixture):
    def test_rows_come_out_of_the_printed_layout(self):
        path = self.pdf(
            [
                ("1", GOOD, "Ali 123456 1403-01-01"),
                ("2", "6100015738451234567", "Reza 654321 1403-01-02"),
            ]
        )
        report = processor.process_file(path, path.name)
        self.assertEqual(csv_rows(report)[1:], ["123456," + GOOD])
        # a PDF has no sheet cells, so the link names the page instead of inventing one
        self.assertIn("صفحهٔ 1", report.problems_csv or "")
        self.assertIn("بارکد", report.problems_csv or "")
        self.assertIn("19 رقم", report.problems_csv or "")  # the length rule, named
        self.assertIn("صفحهٔ 1", (report.problems_csv or "") + report.summary + (report.review_csv or ""))

    def test_a_moved_layout_is_refused_instead_of_returning_nothing(self):
        # The silent version of this bug was an empty tracking.csv: the row exists,
        # the columns are somewhere else, and the warehouse gets "0 ردیف" with a ✅.
        path = self.dir / "shifted.pdf"
        doc = pymupdf.open()
        page = doc.new_page(width=595, height=842)
        for index, (number, barcode, middle) in enumerate([("1", GOOD, "Ali 123456 1403-01-01")]):
            y = 90 + index * 16
            page.insert_text((60, y), number, fontsize=9)
            page.insert_text((120, y), barcode, fontsize=9)
            page.insert_text((300, y), middle, fontsize=9)
        doc.save(path)
        doc.close()
        with self.assertRaises(ValueError) as caught:
            processor.process_file(path, path.name)
        self.assertIn("ساختار PDF شناخته نشد", str(caught.exception))


# ---------------------------------------------------------------------------
# 7.6 the ledger of processed files
# ---------------------------------------------------------------------------


@needs_flow
class TestTrackingLedger(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        tracking_ledger.FILE = Path(self._tmp.name) / "tracking_ledger.json"
        jsonstore_invalidate()

    def test_the_fingerprint_follows_the_bytes_not_the_name(self):
        src = Path(self._tmp.name) / "a.xlsx"
        src.write_bytes(b"one,two\n")
        other = Path(self._tmp.name) / "ب.xlsx"
        other.write_bytes(b"one,two\n")
        grown = Path(self._tmp.name) / "c.xlsx"
        grown.write_bytes(b"one,two\nthree\n")
        self.assertEqual(tracking_ledger.fingerprint(src), tracking_ledger.fingerprint(other))
        self.assertNotEqual(tracking_ledger.fingerprint(src), tracking_ledger.fingerprint(grown))

    def test_remember_then_find_returns_the_earlier_run(self):
        fp = tracking_ledger.fingerprint(self._write("x.xlsx"))
        self.assertIsNone(tracking_ledger.find(fp))
        tracking_ledger.remember(
            fp, user_id=SUDO, fname="x.xlsx", report={"rows": 12, "usable": 10, "errors": 2, "warnings": 1}
        )
        entry = tracking_ledger.find(fp)
        self.assertIsNotNone(entry)
        assert entry is not None
        self.assertEqual((entry["rows"], entry["usable"], entry["errors"]), (12, 10, 2))
        self.assertEqual(entry["user_id"], SUDO)
        describe = tracking_ledger.describe(entry)
        self.assertIn("12 ردیف", describe)
        self.assertIn("❌ 2", describe)

    def test_only_the_counts_are_stored(self):
        tracking_ledger.remember(
            "abc",
            user_id=SUDO,
            fname="orders.xlsx",
            report={"rows": 1, "usable": 1, "errors": 0, "warnings": 0, "file": "orders.xlsx"},
        )
        raw = tracking_ledger.FILE.read_text(encoding="utf-8")
        self.assertNotIn(GOOD, raw)  # no barcodes in the ledger
        self.assertIn('"file": "orders.xlsx"', raw)

    def test_the_ledger_forgets_oldest_first(self):
        for index in range(tracking_ledger.MAX_ENTRIES + 5):
            tracking_ledger.remember(f"fp{index}", user_id=SUDO, fname=f"f{index}.xlsx", report={})
        self.assertEqual(len(tracking_ledger.recent(limit=1000)), tracking_ledger.MAX_ENTRIES)
        self.assertIsNone(tracking_ledger.find("fp0"))
        self.assertIsNotNone(tracking_ledger.find(f"fp{tracking_ledger.MAX_ENTRIES + 4}"))

    def test_a_clear_empties_the_ledger_alone(self):
        tracking_ledger.remember("fp", user_id=SUDO, fname="f.xlsx", report={})
        self.assertEqual(tracking_ledger.clear(), 1)
        self.assertIsNone(tracking_ledger.find("fp"))

    def test_the_start_up_check_prints_the_ledger(self) -> None:
        """``--check-config`` is the only screen a dead ledger can be seen on."""
        import main

        tracking_ledger.remember("fp", user_id=SUDO, fname="f.xlsx", report={"rows": 4, "usable": 4})
        first = io.StringIO()
        with mock.patch.dict(os.environ, {"TISA_DATA_DIR": self._tmp.name}), redirect_stdout(first):
            rc_clean = main.check_config()
        self.assertIn("ledger  :", first.getvalue())
        self.assertIn("1 فایل", first.getvalue(), first.getvalue())
        self.assertEqual(0, rc_clean, first.getvalue())

        # A ledger that cannot be read costs only the repeat warning: said out loud,
        # but it must not turn into "your deployment is broken" either.
        tracking_ledger.FILE.write_text("{not json", encoding="utf-8")
        jsonstore_invalidate()
        second = io.StringIO()
        with mock.patch.dict(os.environ, {"TISA_DATA_DIR": self._tmp.name}), redirect_stdout(second):
            rc_broken = main.check_config()
        self.assertIn("قابل‌خواندن نیست", second.getvalue())
        self.assertEqual(rc_clean, rc_broken, "یک دفتر خراب، پیکربندی را رد نمی‌کند")

    def _write(self, name: str) -> Path:
        path = Path(self._tmp.name) / name
        path.write_bytes(b"header\nrow\n")
        return path


def jsonstore_invalidate() -> None:
    """Drop the shared read cache (the store keys its cache by path)."""
    from bot.services import jsonstore

    jsonstore.invalidate()


# ---------------------------------------------------------------------------
# 7.7 the conversation: buttons that are real handlers, files that are released
# ---------------------------------------------------------------------------


@needs_flow
class TestTrackingFlow(FileFixture):
    def setUp(self) -> None:
        super().setUp()
        self.tmp = Path(self._tmp.name) / "tmp-root"
        self._dirs = (TC.TEMP_DIR, tracking_ledger.FILE)
        TC.TEMP_DIR = self.tmp
        # Outside the temp root on purpose: «the flow cleaned up its download» must not
        # be measured against a ledger file that is meant to survive.
        tracking_ledger.FILE = self.tmp.parent / "tracking_ledger.json"
        jsonstore_invalidate()
        self.addCleanup(self._restore)

    def _restore(self) -> None:
        TC.TEMP_DIR, tracking_ledger.FILE = self._dirs
        jsonstore_invalidate()

    def run_flow(self, update, context):
        import asyncio

        return asyncio.run(TC.on_document(update, context))

    def test_every_button_the_flow_can_render_has_a_handler(self):
        patterns = h.conversation_patterns()
        joined = " ".join(patterns)
        for button in (CB.TRACKING_CONVERT, CB.TRACKING_CANCEL, CB.TRACKING_RETRY, CB.TRACKING_MORE):
            self.assertIn(button.split(":")[-1], joined, f"{button} has no handler")
        # the column answers are a pattern, not a literal, and must be inside the flow
        self.assertTrue(any("trk:pick" in pattern for pattern in patterns), sorted(patterns))

    def test_an_ambiguous_file_is_answered_with_buttons_not_a_guess(self):
        import asyncio

        from telegram.ext import ConversationHandler

        path = self.xlsx(
            [
                ["ردیف", "بارکد", "کد رهگیری", "کد سفارش"],
                [1, GOOD, GOOD2, "123456"],
            ]
        )
        context = h.context()
        update, sent = h.document_update("orders.xlsx", path, user_id=SUDO)
        self.assertEqual(self.run_flow(update, context), TC.MAP_COLUMNS)
        asked = [str(entry[1]) for entry in sent if entry[0] == "text"][-1]
        self.assertIn("کدام ستون؟", asked)
        labels = [button.text for row in update.effective_message.keyboard_rows() for button in row]
        self.assertTrue(any("بارکد" in label for label in labels), labels)
        self.assertTrue(any("کد رهگیری" in label for label in labels), labels)
        # the download is held while the question waits, and it is *this* file
        self.assertTrue(Path(TC.pending_of(context)["path"]).exists())

        # answering with column B («بارکد») writes that column's value
        query_update, _ = h.query_update(f"{TC.PICK_PREFIX}{processor.FIELD_BARCODE}:1", user_id=SUDO)
        query_update.callback_query.message = update.effective_message
        self.assertEqual(asyncio.run(TC.on_pick_column(query_update, context)), TC.ASK_FILE)
        documents = dict(update.effective_message.documents())
        self.assertIn("123456," + GOOD, documents["tracking.csv"].decode("utf-8-sig"))
        # the answer is not remembered as a rule: nothing else changes for the next file
        self.assertEqual(TC.pending_of(context), {})
        self.assertFalse(any(Path(TC.TEMP_DIR).iterdir()))
        self.assertEqual(ConversationHandler.END, -1)  # the flow stays open, not ended

    def test_an_oversized_file_is_refused_before_it_is_downloaded(self) -> None:
        """The size guard is about the download, so the test is too: nothing may land on disk."""
        path = self.xlsx([HEADER, [1, GOOD, "1403-01-01", "امیر ۱۲۳۴۵", "123456", "x"]])
        context = h.context()
        update, sent = h.document_update("big.xlsx", path, user_id=SUDO, file_size=99 * 1024 * 1024)
        self.assertEqual(self.run_flow(update, context), TC.ASK_FILE)
        text = str(sent[-1][1])
        self.assertIn("99.0 MB", text)
        self.assertIn("MAX_FILE_MB", text)
        self.assertFalse(any(Path(TC.TEMP_DIR).iterdir()), "the file was never downloaded")
        self.assertEqual(TC.pending_of(context), {})

    def test_a_repeat_file_is_recognised_and_re_run_on_purpose(self):
        import asyncio

        path = self.xlsx([HEADER, [1, GOOD, "1403-01-01", "امیر ۱۲۳۴۵", "123456", "x"]])
        context = h.context()
        first_update, _ = h.document_update("orders.xlsx", path, user_id=SUDO)
        self.assertEqual(self.run_flow(first_update, context), TC.ASK_FILE)

        repeat_update, sent = h.document_update("orders.xlsx", path, user_id=SUDO)
        self.assertEqual(self.run_flow(repeat_update, context), TC.DUPLICATE_FILE)
        texts = [str(entry[1]) for entry in sent if entry[0] == "text"]
        self.assertTrue(any("قبلاً پردازش شده" in text for text in texts), texts)
        # and no second tracking.csv was produced behind the question
        self.assertEqual(repeat_update.effective_message.documents(), [])

        query_update, _ = h.query_update(CB.TRACKING_RETRY, user_id=SUDO)
        query_update.callback_query.message = repeat_update.effective_message
        self.assertEqual(asyncio.run(TC.on_retry(query_update, context)), TC.ASK_FILE)
        documents = repeat_update.effective_message.documents()
        self.assertEqual([name for name, _ in documents], ["tracking.csv"])

    def test_a_wrong_file_type_is_refused_before_anything_is_downloaded(self):
        path = self.dir / "orders.txt"
        path.write_text("بارکد\n", encoding="utf-8")
        context = h.context()
        update, sent = h.document_update("orders.txt", path, user_id=SUDO)
        self.assertEqual(self.run_flow(update, context), TC.ASK_FILE)
        self.assertIn("فقط فایل‌های xlsx / csv / pdf", str(sent[-1][1]))

    def test_the_instructions_carry_the_limits_the_code_enforces(self):
        with h.patched_settings(h.settings_with(max_rows=1234, max_file_mb=7.5)):
            text = TC.INSTRUCTIONS + "\n" + TC._limits_line()
        self.assertIn("1,234", text)
        self.assertIn("7.5 MB", text)

    def test_leaving_the_flow_releases_the_download(self):
        import asyncio

        from telegram.ext import ConversationHandler

        path = self.xlsx([["ردیف", "بارکد", "کد رهگیری"], [1, GOOD, GOOD2], [2, GOOD2, GOOD]])
        context = h.context()
        update, _ = h.document_update("orders.xlsx", path, user_id=SUDO)
        self.assertEqual(self.run_flow(update, context), TC.MAP_COLUMNS)

        query_update, _ = h.query_update(CB.TRACKING_CANCEL, user_id=SUDO)
        query_update.callback_query.message = update.effective_message
        self.assertEqual(asyncio.run(TC.cb_cancel(query_update, context)), ConversationHandler.END)
        self.assertEqual(TC.pending_of(context), {})
        self.assertFalse(any(Path(TC.TEMP_DIR).iterdir()))

    def test_a_timeout_releases_the_download_too(self):
        import asyncio

        path = self.xlsx([["ردیف", "بارکد", "کد رهگیری"], [1, GOOD, GOOD2]])
        context = h.context()
        update, _ = h.document_update("orders.xlsx", path, user_id=SUDO)
        self.assertEqual(self.run_flow(update, context), TC.MAP_COLUMNS)
        asyncio.run(TC.on_timeout(update, context))
        self.assertEqual(TC.pending_of(context), {})

    def test_a_non_authorized_user_cannot_open_the_flow(self):
        import asyncio

        from telegram.ext import ConversationHandler

        context = h.context()
        update, update_seen = h.query_update(CB.TRACKING_CONVERT, user_id=999)
        state = asyncio.run(TC.entry(update, context))
        self.assertEqual(state, ConversationHandler.END)
        self.assertIn("⛔ دسترسی ندارید.", answers(update_seen))


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
